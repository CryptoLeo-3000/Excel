import openpyxl as xl
from openpyxl.chart import BarChart,Reference

def charts(filename):
    wb = xl.load_workbook(f'Excel/{filename}.xlsx')
    sheet = wb['Sheet1']

    values = Reference(sheet, min_row=2, max_row=sheet.max_row, min_col=4, max_col=4)
    chart = BarChart()
    chart.add_data(values)
    sheet.add_chart(chart, 'e2')

    wb.save(f'Excel/{filename}.xlsx')